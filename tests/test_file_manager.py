"""Tests for file manager safety guarantees."""

import os
import tempfile
import time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_mcp_server import add_allowed_directory
from excel_mcp_server.config import settings
from excel_mcp_server.utils.file_manager import FileManager
from excel_mcp_server.utils.security import SecurityError


def _create_workbook(file_path: Path, value: str = "before") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = value
    workbook.save(file_path)
    workbook.close()


def test_safe_write_uses_atomic_temp_file() -> None:
    """Workbook saves should go through a temp file before replacing the target."""
    with tempfile.TemporaryDirectory() as directory:
        add_allowed_directory(directory)
        file_path = Path(directory) / "atomic.xlsx"
        _create_workbook(file_path)

        manager = FileManager()
        save_targets: list[str] = []

        with manager.safe_write(file_path) as workbook:
            workbook.active["A1"] = "after"
            original_save = workbook.save

            def wrapped_save(target: str) -> None:
                save_targets.append(target)
                original_save(target)

            workbook.save = wrapped_save  # type: ignore[method-assign]

        assert save_targets
        assert Path(save_targets[0]) != file_path
        assert Path(save_targets[0]).parent == file_path.parent

        reloaded = load_workbook(file_path, data_only=True)
        assert reloaded.active["A1"].value == "after"
        reloaded.close()


def test_safe_open_cleans_stale_lock_files() -> None:
    """Stale lock files should be removed automatically before opening a workbook."""
    with tempfile.TemporaryDirectory() as directory:
        add_allowed_directory(directory)
        file_path = Path(directory) / "locked.xlsx"
        _create_workbook(file_path)

        lock_dir = file_path.parent / ".excel_locks"
        lock_dir.mkdir(exist_ok=True)
        stale_lock = lock_dir / "orphan.lock"
        stale_lock.write_text("stale")

        stale_at = time.time() - 3600
        os.utime(stale_lock, (stale_at, stale_at))

        manager = FileManager()
        with manager.safe_open(file_path, read_only=True):
            pass

        assert not stale_lock.exists()
        assert not (lock_dir / f"{file_path.name}.lock").exists()


def test_safe_open_rejects_allowed_directory_prefix_bypass() -> None:
    """Paths that merely share an allowed-directory prefix must be rejected."""
    original_allowed_directories = list(settings.security.allowed_directories)
    settings.security.allowed_directories.clear()

    try:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            allowed_dir = root / "allowed"
            bypass_dir = root / "allowed_evil"
            allowed_dir.mkdir()
            bypass_dir.mkdir()
            add_allowed_directory(allowed_dir)

            file_path = bypass_dir / "escape.xlsx"
            _create_workbook(file_path)

            manager = FileManager()

            with pytest.raises(SecurityError, match="allowed directories"):
                with manager.safe_open(file_path, read_only=True):
                    pass

            assert not (bypass_dir / ".excel_locks").exists()
    finally:
        settings.security.allowed_directories[:] = original_allowed_directories
