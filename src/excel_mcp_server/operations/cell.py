"""Cell operations"""

import logging
from pathlib import Path

from openpyxl.utils import get_column_letter

from ..models import (
    CellReadRequest,
    CellWriteRequest,
    CellResult,
    FormulaWriteRequest,
    RangeReadRequest,
    RangeResult,
    RangeWriteRequest,
)
from ..utils.file_manager import file_manager

logger = logging.getLogger(__name__)


def read_cell(request: CellReadRequest) -> CellResult:
    """Read a value from a cell."""
    try:
        path = Path(request.file_path).resolve()
        
        with file_manager.safe_open(path, read_only=True, data_only=True) as wb:
            if request.sheet_name not in wb.sheetnames:
                return CellResult(
                    success=False,
                    message=f"Sheet '{request.sheet_name}' not found. Available: {wb.sheetnames}",
                    error_code="SHEET_NOT_FOUND"
                )
            
            ws = wb[request.sheet_name]
            value = ws[request.cell].value
        
        return CellResult(
            success=True,
            message=f"Value read from {request.cell}",
            cell=request.cell,
            value=value
        )
        
    except Exception as e:
        logger.exception(f"Failed to read cell: {request.cell}")
        return CellResult(
            success=False,
            message=f"Failed to read cell: {str(e)}",
            error_code="READ_ERROR"
        )


def write_cell(request: CellWriteRequest) -> CellResult:
    """Write a value to a cell."""
    try:
        path = Path(request.file_path).resolve()
        
        with file_manager.safe_write(path) as wb:
            if request.sheet_name not in wb.sheetnames:
                return CellResult(
                    success=False,
                    message=f"Sheet '{request.sheet_name}' not found. Available: {wb.sheetnames}",
                    error_code="SHEET_NOT_FOUND"
                )
            
            ws = wb[request.sheet_name]
            ws[request.cell] = request.value
        
        return CellResult(
            success=True,
            message=f"Value written to {request.cell}",
            cell=request.cell,
            value=request.value
        )
        
    except Exception as e:
        logger.exception(f"Failed to write cell: {request.cell}")
        return CellResult(
            success=False,
            message=f"Failed to write cell: {str(e)}",
            error_code="WRITE_ERROR"
        )


def read_range(request: RangeReadRequest) -> RangeResult:
    """Read data from a range of cells."""
    try:
        path = Path(request.file_path).resolve()
        
        with file_manager.safe_open(path, read_only=True, data_only=True) as wb:
            if request.sheet_name not in wb.sheetnames:
                return RangeResult(
                    success=False,
                    message=f"Sheet '{request.sheet_name}' not found. Available: {wb.sheetnames}",
                    error_code="SHEET_NOT_FOUND"
                )
            
            ws = wb[request.sheet_name]
            cell_range = ws[request.range_ref]
            
            data = []
            if isinstance(cell_range, tuple):
                for row in cell_range:
                    if isinstance(row, tuple):
                        data.append([cell.value for cell in row])
                    else:
                        data.append([row.value])
            else:
                data = [[cell_range.value]]
        
        rows = len(data)
        cols = len(data[0]) if data else 0
        
        return RangeResult(
            success=True,
            message=f"Data read from range {request.range_ref}",
            range_ref=request.range_ref,
            rows=rows,
            cols=cols,
            data=data
        )
        
    except Exception as e:
        logger.exception(f"Failed to read range: {request.range_ref}")
        return RangeResult(
            success=False,
            message=f"Failed to read range: {str(e)}",
            error_code="READ_ERROR"
        )


def write_range(request: RangeWriteRequest) -> RangeResult:
    """Write data to a range of cells."""
    try:
        path = Path(request.file_path).resolve()
        
        # Parse start cell
        import re
        match = re.match(r'^([A-Z]+)(\d+)$', request.start_cell, re.IGNORECASE)
        if not match:
            return RangeResult(
                success=False,
                message=f"Invalid start cell: {request.start_cell}",
                error_code="INVALID_CELL"
            )
        
        start_col_str, start_row_str = match.groups()
        start_col = 0
        for char in start_col_str.upper():
            start_col = start_col * 26 + (ord(char) - ord('A') + 1)
        start_row = int(start_row_str)
        
        with file_manager.safe_write(path) as wb:
            if request.sheet_name not in wb.sheetnames:
                return RangeResult(
                    success=False,
                    message=f"Sheet '{request.sheet_name}' not found. Available: {wb.sheetnames}",
                    error_code="SHEET_NOT_FOUND"
                )
            
            ws = wb[request.sheet_name]
            
            rows_written = 0
            cols_written = 0
            
            for row_idx, row_data in enumerate(request.data):
                for col_idx, value in enumerate(row_data):
                    cell = ws.cell(
                        row=start_row + row_idx,
                        column=start_col + col_idx,
                        value=value
                    )
                    cols_written = max(cols_written, col_idx + 1)
                rows_written += 1
        
        return RangeResult(
            success=True,
            message=f"Data written to range starting at {request.start_cell}",
            range_ref=request.start_cell,
            rows=rows_written,
            cols=cols_written
        )
        
    except Exception as e:
        logger.exception(f"Failed to write range: {request.start_cell}")
        return RangeResult(
            success=False,
            message=f"Failed to write range: {str(e)}",
            error_code="WRITE_ERROR"
        )


def write_formula(request: FormulaWriteRequest) -> CellResult:
    """Write a formula to a cell."""
    try:
        path = Path(request.file_path).resolve()
        formula = request.formula if request.formula.startswith('=') else f'={request.formula}'
        
        with file_manager.safe_write(path) as wb:
            if request.sheet_name not in wb.sheetnames:
                return CellResult(
                    success=False,
                    message=f"Sheet '{request.sheet_name}' not found. Available: {wb.sheetnames}",
                    error_code="SHEET_NOT_FOUND"
                )
            
            ws = wb[request.sheet_name]
            ws[request.cell] = formula
        
        return CellResult(
            success=True,
            message=f"Formula written to {request.cell}",
            cell=request.cell,
            value=formula
        )
        
    except Exception as e:
        logger.exception(f"Failed to write formula: {request.cell}")
        return CellResult(
            success=False,
            message=f"Failed to write formula: {str(e)}",
            error_code="WRITE_ERROR"
        )
