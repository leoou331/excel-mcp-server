"""Formatting operations"""

import logging
from pathlib import Path

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from ..models import FontFormatRequest, FillFormatRequest, NumberFormatRequest, RangeResult
from ..utils.file_manager import file_manager

logger = logging.getLogger(__name__)


def format_font(request: FontFormatRequest) -> RangeResult:
    """Apply font formatting to a range."""
    try:
        path = Path(request.file_path).resolve()
        
        font_kwargs = {}
        if request.bold is not None:
            font_kwargs['bold'] = request.bold
        if request.italic is not None:
            font_kwargs['italic'] = request.italic
        if request.font_size is not None:
            font_kwargs['size'] = request.font_size
        if request.color is not None:
            font_kwargs['color'] = request.color
        
        font = Font(**font_kwargs) if font_kwargs else None
        
        with file_manager.safe_write(path) as wb:
            if request.sheet_name not in wb.sheetnames:
                return RangeResult(
                    success=False,
                    message=f"Sheet '{request.sheet_name}' not found",
                    error_code="SHEET_NOT_FOUND"
                )
            
            ws = wb[request.sheet_name]
            
            for row in ws[request.range_ref]:
                if isinstance(row, tuple):
                    for cell in row:
                        if font:
                            cell.font = font
                else:
                    if font:
                        row.font = font
        
        return RangeResult(
            success=True,
            message=f"Font formatting applied to {request.range_ref}",
            range_ref=request.range_ref
        )
        
    except Exception as e:
        logger.exception(f"Failed to format font: {request.range_ref}")
        return RangeResult(
            success=False,
            message=f"Failed to format font: {str(e)}",
            error_code="FORMAT_ERROR"
        )


def format_fill(request: FillFormatRequest) -> RangeResult:
    """Apply fill (background color) formatting to a range."""
    try:
        path = Path(request.file_path).resolve()
        
        fill = PatternFill(
            start_color=request.color,
            end_color=request.color,
            fill_type='solid'
        )
        
        with file_manager.safe_write(path) as wb:
            if request.sheet_name not in wb.sheetnames:
                return RangeResult(
                    success=False,
                    message=f"Sheet '{request.sheet_name}' not found",
                    error_code="SHEET_NOT_FOUND"
                )
            
            ws = wb[request.sheet_name]
            
            for row in ws[request.range_ref]:
                if isinstance(row, tuple):
                    for cell in row:
                        cell.fill = fill
                else:
                    row.fill = fill
        
        return RangeResult(
            success=True,
            message=f"Fill formatting applied to {request.range_ref}",
            range_ref=request.range_ref
        )
        
    except Exception as e:
        logger.exception(f"Failed to format fill: {request.range_ref}")
        return RangeResult(
            success=False,
            message=f"Failed to format fill: {str(e)}",
            error_code="FORMAT_ERROR"
        )


def format_number(request: NumberFormatRequest) -> RangeResult:
    """Apply number formatting to a range."""
    try:
        path = Path(request.file_path).resolve()
        
        with file_manager.safe_write(path) as wb:
            if request.sheet_name not in wb.sheetnames:
                return RangeResult(
                    success=False,
                    message=f"Sheet '{request.sheet_name}' not found",
                    error_code="SHEET_NOT_FOUND"
                )
            
            ws = wb[request.sheet_name]
            
            for row in ws[request.range_ref]:
                if isinstance(row, tuple):
                    for cell in row:
                        cell.number_format = request.format_string
                else:
                    row.number_format = request.format_string
        
        return RangeResult(
            success=True,
            message=f"Number formatting applied to {request.range_ref}",
            range_ref=request.range_ref
        )
        
    except Exception as e:
        logger.exception(f"Failed to format number: {request.range_ref}")
        return RangeResult(
            success=False,
            message=f"Failed to format number: {str(e)}",
            error_code="FORMAT_ERROR"
        )
