"""Safe file management with locking and context managers"""

import logging
import os
import tempfile
import time
from contextlib import contextmanager
from pathlib import Path
from typing import Generator, Optional

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook
from openpyxl.workbook import Workbook as WorkbookType

from ..config import settings
from .security import SecurityError

logger = logging.getLogger(__name__)


class FileManager:
    """Manages Excel file operations with safety guarantees."""

    LOCK_DIRNAME = ".excel_locks"
    LOCK_SUFFIX = ".lock"

    def _get_lock_dir(self, path: Path) -> Path:
        """Return the directory used for workbook lock files."""
        return path.parent / self.LOCK_DIRNAME

    def _get_lock_path(self, path: Path) -> Path:
        """Return the lock file path for a workbook."""
        return self._get_lock_dir(path) / f"{path.name}{self.LOCK_SUFFIX}"

    def _cleanup_lock_file(self, lock_path: Path, stale_only: bool = False) -> None:
        """Best-effort cleanup for orphaned lock files."""
        if not lock_path.exists():
            return

        if stale_only:
            try:
                age_seconds = time.time() - lock_path.stat().st_mtime
            except OSError as exc:
                logger.debug("Unable to inspect lock file %s: %s", lock_path, exc)
                return

            if age_seconds < settings.security.lock_stale_seconds:
                return

        cleanup_lock = FileLock(str(lock_path), timeout=0)
        try:
            cleanup_lock.acquire(blocking=False)
        except Timeout:
            return
        except Exception as exc:
            logger.debug("Unable to probe lock file %s: %s", lock_path, exc)
            return

        try:
            lock_path.unlink(missing_ok=True)
            logger.debug("Removed orphaned lock file: %s", lock_path)
        except OSError as exc:
            logger.debug("Unable to remove lock file %s: %s", lock_path, exc)
        finally:
            try:
                cleanup_lock.release()
            except Exception as exc:
                logger.debug("Unable to release cleanup probe for %s: %s", lock_path, exc)

    def _cleanup_stale_locks(self, lock_dir: Path) -> None:
        """Remove stale lock files before taking a new lock."""
        if not lock_dir.exists():
            return

        for lock_path in lock_dir.glob(f"*{self.LOCK_SUFFIX}"):
            self._cleanup_lock_file(lock_path, stale_only=True)

    def _save_workbook_atomically(self, workbook: WorkbookType, path: Path) -> None:
        """Persist workbook data via a temp file and atomic replace."""
        temp_fd, temp_name = tempfile.mkstemp(
            prefix=f".{path.stem}.",
            suffix=path.suffix,
            dir=path.parent,
        )
        os.close(temp_fd)
        temp_path = Path(temp_name)

        try:
            workbook.save(str(temp_path))
            os.replace(temp_path, path)
        except Exception:
            temp_path.unlink(missing_ok=True)
            raise
    
    @contextmanager
    def safe_open(
        self,
        file_path: str | Path,
        read_only: bool = False,
        data_only: bool = False,
        timeout: Optional[float] = None,
    ) -> Generator[WorkbookType, None, None]:
        """Safely open an Excel workbook with file locking."""
        path = Path(file_path).resolve()
        timeout = timeout or settings.security.lock_timeout_seconds
        
        # Create lock file path
        lock_dir = self._get_lock_dir(path)
        lock_dir.mkdir(exist_ok=True)
        self._cleanup_stale_locks(lock_dir)
        lock_path = self._get_lock_path(path)
        lock = FileLock(str(lock_path), timeout=timeout)
        
        wb: Optional[WorkbookType] = None
        lock_acquired = False
        
        try:
            lock.acquire(blocking=True)
            lock_acquired = True
            logger.debug(f"Lock acquired for {path}")
            
            if path.exists():
                wb = load_workbook(
                    str(path),
                    read_only=read_only,
                    data_only=data_only,
                    keep_links=False,
                )
            elif not read_only:
                wb = Workbook()
            else:
                raise SecurityError(f"File not found: {path}")
            
            yield wb
            
        except Timeout:
            logger.error(f"Failed to acquire lock for {path}")
            raise SecurityError(f"File is locked by another process: {path}")
        except Exception as e:
            logger.exception(f"Error opening workbook {path}")
            raise
        finally:
            if wb is not None:
                try:
                    wb.close()
                    logger.debug(f"Workbook closed: {path}")
                except Exception as e:
                    logger.warning(f"Error closing workbook: {e}")
            
            if lock_acquired:
                try:
                    lock.release()
                    logger.debug(f"Lock released for {path}")
                except Exception as e:
                    logger.warning(f"Error releasing lock: {e}")
                self._cleanup_lock_file(lock_path)
    
    @contextmanager
    def safe_write(
        self,
        file_path: str | Path,
        timeout: Optional[float] = None,
    ) -> Generator[WorkbookType, None, None]:
        """Safely open an Excel workbook for writing with auto-save."""
        path = Path(file_path).resolve()
        
        with self.safe_open(path, read_only=False, data_only=False, timeout=timeout) as wb:
            try:
                yield wb
                self._save_workbook_atomically(wb, path)
                logger.info(f"Workbook saved: {path}")
            except Exception as e:
                logger.error(f"Error during write, changes not saved: {e}")
                raise
    
    def safe_create(
        self,
        file_path: str | Path,
        timeout: Optional[float] = None,
    ) -> bool:
        """Create a new Excel workbook. Returns True if created."""
        path = Path(file_path).resolve()
        
        if path.exists():
            return False
        
        with self.safe_write(path, timeout=timeout) as wb:
            pass  # Workbook created automatically
        
        return True


file_manager = FileManager()
